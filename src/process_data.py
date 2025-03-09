import fitz  # PyMuPDF do PDF
import pandas as pd  # Pandas do Excela
import docx  # Obsługa plików DOCX
from openpyxl import load_workbook  # Excel

def extract_text_from_pdf(pdf_path):
    """Ekstrakcja tekstu z pliku PDF."""
    doc = fitz.open(pdf_path)
    text = "\n".join([page.get_text() for page in doc])
    return text

def extract_text_from_docx(docx_path):
    """Ekstrakcja tekstu z pliku DOCX."""
    doc = docx.Document(docx_path)
    text = "\n".join([para.text for para in doc.paragraphs])
    return text

def extract_data_from_xlsx(xlsx_path):
    """Ekstrakcja danych z pliku Excel (XLSX)."""
    wb = load_workbook(xlsx_path)
    sheet = wb.active
    data = pd.DataFrame(sheet.values)
    return data.to_string()

# Testowy kod - można usunąć później
if __name__ == "__main__":
    print("Moduł do przetwarzania dokumentów załadowany!")
